# -*- coding: utf-8 -*-
"""
真实数据导入工具：把 Excel/CSV 业务数据导入 data/shop.db。
  - 按 SKU/订单号 UPSERT（重复导入=更新，不产生重复行）
  - 中文列头自动映射，缺必填列/非法值时明确报错
  - --dry-run 只校验不写入；--with-items 同时导入订单明细 sheet
  - --skip-sample 跳过"示例："开头的行

用法：
  python data/import_real_data.py 商品表.xlsx
  python data/import_real_data.py 订单表.xlsx --with-items
  python data/import_real_data.py 商品表.xlsx --dry-run
模板生成：python data/make_templates.py
"""
import argparse
import csv
import json
import sqlite3
import sys
from contextlib import closing
from datetime import datetime
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

DB_PATH = Path(__file__).resolve().parent / "shop.db"

# 列头映射：Excel 中文列名 → 数据库字段（带 * 为模板必填列名，同时兼容常见简称）
PRODUCT_COLS = {
    "SKU编号*": "sku", "SKU编号": "sku", "SKU": "sku", "sku": "sku",
    "商品名称*": "name", "商品名称": "name", "名称": "name",
    "类目": "category", "分类": "category",
    "价格(元)*": "price", "价格": "price", "售价": "price",
    "库存数量*": "stock", "库存": "stock",
    "标签(分号分隔)": "tags", "标签": "tags",
    "评分(0-5)": "rating", "评分": "rating",
    "状态(在售/下架)": "status", "状态": "status",
    "商品描述": "description", "描述": "description",
}
ORDER_COLS = {
    "订单号*": "order_no", "订单号": "order_no",
    "顾客姓名*": "customer_name", "顾客姓名": "customer_name",
    "顾客手机号": "phone", "手机号": "phone",
    "订单状态*": "status", "订单状态": "status",
    "订单总额(元)*": "total_amount", "订单总额": "total_amount", "总额": "total_amount",
    "下单时间": "created_at", "支付时间": "payment_time", "发货时间": "ship_time",
    "快递公司": "carrier", "运单号": "tracking_no",
    "物流状态": "logistics_status", "物流最后更新时间": "logistics_updated_at",
    "物流备注": "logistics_note", "签收时间": "receipt_time",
    "售后备注": "after_sale_note", "收货地址": "address",
}
ITEM_COLS = {
    "订单号*": "order_no", "订单号": "order_no",
    "SKU编号*": "sku", "SKU编号": "sku", "SKU": "sku",
    "数量*": "qty", "数量": "qty",
    "单价(元)*": "unit_price", "单价": "unit_price",
    "明细状态(正常/已退款/售后中)": "item_status", "明细状态": "item_status",
}
ALL_COL_MAPS = [PRODUCT_COLS, ORDER_COLS, ITEM_COLS]
VALID_ORDER_STATUS = {"待付款", "待发货", "已发货", "已签收", "售后中", "已退款", "已取消"}
TIME_FORMATS = ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y-%m-%d")


# ---------------------------------------------------------------- 读取
def load_rows(path: Path, sheet: str | None) -> list[dict]:
    """读取 xlsx/csv → dict 列表（键=数据库字段，按列头映射表自动对应）。"""
    if path.suffix.lower() == ".csv":
        with open(path, encoding="utf-8-sig") as f:
            raw = list(csv.reader(f))
        headers, data = raw[0], raw[1:]
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        if sheet and sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            # 跳过【填写说明】等说明页，取第一个数据 sheet
            ws = next((wb[n] for n in wb.sheetnames if not n.startswith("【")),
                      wb[wb.sheetnames[0]])
        rows = [[cell for cell in r] for r in ws.iter_rows(values_only=True)]
        if not rows:
            return []
        headers, data = list(rows[0]), rows[1:]

    # 建立映射：数据库字段 → 列下标（多个别名列按先到先得）
    field_idx: dict[str, int] = {}
    for col_map in ALL_COL_MAPS:
        for cn, field in col_map.items():
            if cn in headers and field not in field_idx:
                field_idx[field] = headers.index(cn)

    out = []
    for r in data:
        item = {f: (r[i] if i < len(r) else None) for f, i in field_idx.items()}
        if any(v is not None and str(v).strip() != "" for v in item.values()):
            out.append(item)
    return out


def cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    return str(v).strip()


def fmt_time(v) -> str | None:
    """统一时间格式（售后政策按 %Y-%m-%d %H:%M 解析）。"""
    s = cell(v)
    if not s:
        return None
    for fmt in TIME_FORMATS:
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d %H:%M")
        except ValueError:
            continue
    return s  # 无法解析时保留原样，便于人工发现


def _is_num(s: str) -> bool:
    try:
        float(s)
        return True
    except ValueError:
        return False


# ---------------------------------------------------------------- 商品导入
def import_products(rows: list[dict], dry: bool, skip_sample: bool):
    ok, skipped = 0, 0
    with closing(sqlite3.connect(DB_PATH)) as conn:
        for i, r in enumerate(rows, start=2):
            sku, name = cell(r.get("sku")), cell(r.get("name"))
            if skip_sample and (sku.startswith("示例") or name.startswith("示例")):
                skipped += 1
                continue
            if not sku or not name:
                raise SystemExit(f"[商品表 第{i}行] 缺少必填列：SKU编号/商品名称")
            price = cell(r.get("price")) or "0"
            stock = cell(r.get("stock")) or "0"
            if not _is_num(price) or not stock.isdigit():
                raise SystemExit(
                    f"[商品表 第{i}行] 价格/库存必须是数字：价格={price} 库存={stock}")
            tags_raw = cell(r.get("tags")).replace("；", ";").replace("，", ";") \
                                              .replace(",", ";")
            tag_list = [t.strip() for t in tags_raw.split(";") if t.strip()]
            conn.execute(
                "INSERT INTO products (sku, name, category, price, stock, tags, "
                "rating, status, description) VALUES (?,?,?,?,?,?,?,?,?) "
                "ON CONFLICT(sku) DO UPDATE SET name=excluded.name, "
                "category=excluded.category, price=excluded.price, "
                "stock=excluded.stock, tags=excluded.tags, rating=excluded.rating, "
                "status=excluded.status, description=excluded.description",
                (sku, name, cell(r.get("category")) or "未分类",
                 float(price), int(stock),
                 json.dumps(tag_list, ensure_ascii=False),
                 float(cell(r.get("rating")) or 0),
                 cell(r.get("status")) or "在售", cell(r.get("description"))))
            ok += 1
        if dry:
            conn.rollback()
        else:
            conn.commit()
    return ok, skipped


# ---------------------------------------------------------------- 订单导入
def import_orders(rows: list[dict], item_rows: list[dict], dry: bool,
                  skip_sample: bool):
    ok, items_ok, skipped = 0, 0, 0
    with closing(sqlite3.connect(DB_PATH)) as conn:
        for i, r in enumerate(rows, start=2):
            order_no = cell(r.get("order_no"))
            name = cell(r.get("customer_name"))
            if skip_sample and (order_no.startswith("示例") or name.startswith("示例")):
                skipped += 1
                continue
            if not order_no or not name:
                raise SystemExit(f"[订单表 第{i}行] 缺少必填列：订单号/顾客姓名")
            status = cell(r.get("status")) or "已签收"
            if status not in VALID_ORDER_STATUS:
                raise SystemExit(
                    f"[订单表 第{i}行] 订单状态非法: '{status}'，"
                    f"必须是 {'/'.join(sorted(VALID_ORDER_STATUS))}")
            total = cell(r.get("total_amount")) or "0"
            if not _is_num(total):
                raise SystemExit(f"[订单表 第{i}行] 订单总额必须是数字: {total}")

            # 顾客：同姓名+手机号复用已有档案，否则新建 U 编号顾客
            phone = cell(r.get("phone")) or None
            found = conn.execute(
                "SELECT customer_id FROM customers WHERE name=? "
                "AND (phone=? OR ? IS NULL)", (name, phone, phone)).fetchone()
            if found:
                cid = found[0]
            else:
                n = conn.execute(
                    "SELECT COUNT(*) FROM customers WHERE customer_id "
                    "LIKE 'U%'").fetchone()[0]
                cid = f"U{n + 1:04d}"
                conn.execute(
                    "INSERT INTO customers (customer_id, name, phone, vip_level, "
                    "points, created_at) VALUES (?,?,?,0,0,?)",
                    (cid, name, phone, fmt_time(r.get("created_at"))
                     or datetime.now().strftime("%Y-%m-%d %H:%M")))

            conn.execute(
                "INSERT INTO orders (order_no, customer_id, status, total_amount, "
                "payment_time, ship_time, carrier, tracking_no, logistics_status, "
                "logistics_updated_at, logistics_note, receipt_time, "
                "after_sale_note, receiver_name, receiver_phone, address, "
                "created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) "
                "ON CONFLICT(order_no) DO UPDATE SET status=excluded.status, "
                "total_amount=excluded.total_amount, "
                "payment_time=excluded.payment_time, ship_time=excluded.ship_time, "
                "carrier=excluded.carrier, tracking_no=excluded.tracking_no, "
                "logistics_status=excluded.logistics_status, "
                "logistics_updated_at=excluded.logistics_updated_at, "
                "logistics_note=excluded.logistics_note, "
                "receipt_time=excluded.receipt_time, "
                "after_sale_note=excluded.after_sale_note, "
                "receiver_name=excluded.receiver_name, "
                "receiver_phone=excluded.receiver_phone, "
                "address=excluded.address, created_at=excluded.created_at",
                (order_no, cid, status, float(total),
                 fmt_time(r.get("payment_time")), fmt_time(r.get("ship_time")),
                 cell(r.get("carrier")), cell(r.get("tracking_no")),
                 cell(r.get("logistics_status")),
                 fmt_time(r.get("logistics_updated_at")),
                 cell(r.get("logistics_note")), fmt_time(r.get("receipt_time")),
                 cell(r.get("after_sale_note")), name, phone,
                 cell(r.get("address")),
                 fmt_time(r.get("created_at"))
                 or datetime.now().strftime("%Y-%m-%d %H:%M")))
            ok += 1

        for i, r in enumerate(item_rows, start=2):
            order_no, sku = cell(r.get("order_no")), cell(r.get("sku"))
            qty, price = cell(r.get("qty")), cell(r.get("unit_price"))
            if not order_no or not sku:
                continue
            if skip_sample and (order_no.startswith("示例") or sku.startswith("示例")):
                continue
            if not conn.execute("SELECT 1 FROM orders WHERE order_no=?",
                                (order_no,)).fetchone():
                raise SystemExit(
                    f"[订单明细 第{i}行] 订单号 {order_no} 不存在——请先导入订单主表")
            if not conn.execute("SELECT 1 FROM products WHERE sku=?",
                                (sku,)).fetchone():
                raise SystemExit(
                    f"[订单明细 第{i}行] SKU {sku} 不在商品表中——请先导入商品表")
            if not qty.isdigit() or not _is_num(price):
                raise SystemExit(
                    f"[订单明细 第{i}行] 数量/单价必须是数字: {qty}/{price}")
            conn.execute("DELETE FROM order_items WHERE order_no=? AND sku=?",
                         (order_no, sku))
            conn.execute(
                "INSERT INTO order_items (order_no, sku, qty, unit_price, "
                "item_status) VALUES (?,?,?,?,?)",
                (order_no, sku, int(qty), float(price),
                 cell(r.get("item_status")) or "正常"))
            items_ok += 1

        if dry:
            conn.rollback()
        else:
            conn.commit()
    return ok, items_ok, skipped


def main():
    ap = argparse.ArgumentParser(description="真实数据导入工具")
    ap.add_argument("file", help="Excel(.xlsx) 或 CSV 文件路径")
    ap.add_argument("--with-items", action="store_true",
                    help="同时导入『订单明细』sheet（订单表用）")
    ap.add_argument("--dry-run", action="store_true", help="只校验，不写库")
    ap.add_argument("--skip-sample", action="store_true",
                    help="跳过以『示例』开头的行")
    ap.add_argument("--sheet", default=None, help="指定 sheet 名（默认第一个）")
    args = ap.parse_args()

    path = Path(args.file)
    if not path.exists():
        raise SystemExit(f"文件不存在: {path}")
    kind = "orders" if "订单" in path.name else "products"
    print(f"导入目标: {kind} | 文件: {path.name} | "
          f"模式: {'DRY-RUN(只校验)' if args.dry_run else '写入'}")

    if kind == "products":
        rows = load_rows(path, args.sheet)
        if not rows:
            raise SystemExit("未读到有效数据行：请确认文件里有数据 sheet"
                             "（默认跳过【填写说明】页），或用 --sheet 指定。")
        ok, skipped = import_products(rows, args.dry_run, args.skip_sample)
        print(f"商品导入: 成功 {ok} 行, 跳过 {skipped} 行")
    else:
        rows = load_rows(path, args.sheet)
        if not rows:
            raise SystemExit("未读到有效数据行：请确认文件里有数据 sheet"
                             "（默认跳过【填写说明】页），或用 --sheet 指定。")
        item_rows = load_rows(path, "订单明细") if args.with_items else []
        if args.with_items and not item_rows:
            print("[警告] 未找到『订单明细』sheet，只导入订单主表")
        ok, items_ok, skipped = import_orders(rows, item_rows, args.dry_run,
                                              args.skip_sample)
        print(f"订单导入: 主表 {ok} 行, 明细 {items_ok} 行, 跳过 {skipped} 行")

    if not args.dry_run:
        with closing(sqlite3.connect(DB_PATH)) as conn:
            for t in ("products", "orders", "order_items", "customers"):
                n = conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
                print(f"  {t:<12} 当前共 {n} 行")
        print("完成！去聊天页问 AI 新导入的商品/订单验证效果。")
    else:
        print("DRY-RUN 校验通过。去掉 --dry-run 重新执行即可真正写入。")


if __name__ == "__main__":
    main()
