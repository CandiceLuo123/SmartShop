# -*- coding: utf-8 -*-
"""
生成真实数据导入模板 Excel（data/templates/ 下）：
  商品表模板.xlsx   —— products 表
  订单表模板.xlsx   —— orders + 订单明细
每个模板含：【填写说明】sheet + 示例行（导入时可跳过或删除）。
用法：python data/make_templates.py
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
OUT = Path(__file__).resolve().parent / "templates"

HEAD_FILL = PatternFill("solid", fgColor="1A73E8")
HEAD_FONT = Font(color="FFFFFF", bold=True)
NOTE_FILL = PatternFill("solid", fgColor="FEF7E0")


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = 22
    ws.freeze_panes = "A2"


def add_note_sheet(wb, title, lines):
    ws = wb.create_sheet("【填写说明】", 0)
    ws.column_dimensions["A"].width = 100
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    for i, line in enumerate(lines, start=3):
        c = ws.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if line.startswith("※"):
            c.fill = NOTE_FILL
    return ws


def make_products():
    wb = Workbook()
    ws = wb.active
    ws.title = "商品表"
    headers = ["SKU编号*", "商品名称*", "类目", "价格(元)*", "库存数量*",
               "标签(分号分隔)", "评分(0-5)", "状态(在售/下架)", "商品描述"]
    ws.append(headers)
    ws.append(["SKU-9001", "示例：无线蓝牙耳机 Pro", "音频", 299, 50,
               "降噪;长续航", 4.8, "在售", "示例行：导入前请删除本行"])
    style_header(ws, len(headers))
    add_note_sheet(wb, "商品表模板 —— 导入到店铺数据库 products 表", [
        "1. 带星号(*)的是必填列，其余可留空。",
        "2. SKU编号：商品的唯一编码，重复导入时按 SKU 覆盖更新（不会产生重复）。",
        "3. 标签：用英文分号 ; 分隔，如：降噪;运动;IPX5防水。标签会影响 AI 推荐。",
        "4. 状态：在售 / 下架。下架商品不会被推荐。",
        "5. 评分：0-5 的数字，影响推荐排序。",
        "※ 第一数据行是示例，导入前请删除，或使用 --skip-sample 参数跳过。",
        "※ 保存为 .xlsx 后执行：python data/import_real_data.py 你的文件.xlsx",
    ])
    wb.save(OUT / "商品表模板.xlsx")


def make_orders():
    wb = Workbook()
    ws = wb.active
    ws.title = "订单表"
    headers = ["订单号*", "顾客姓名*", "顾客手机号", "订单状态*",
               "订单总额(元)*", "下单时间", "支付时间", "发货时间",
               "快递公司", "运单号", "物流状态", "物流最后更新时间",
               "物流备注", "签收时间", "售后备注", "收货地址"]
    ws.append(headers)
    ws.append(["SO20261001001", "示例：张三", "138****0000", "已签收", 318,
               "2026-10-01 10:00", "2026-10-01 10:05", "2026-10-02 09:00",
               "中通", "ZT000111222333", "已签收", "2026-10-04 14:00",
               "本人签收", "2026-10-04 14:00", "", "示例行：导入前请删除本行"])
    style_header(ws, len(headers))

    ws2 = wb.create_sheet("订单明细")
    headers2 = ["订单号*", "SKU编号*", "数量*", "单价(元)*", "明细状态(正常/已退款/售后中)"]
    ws2.append(headers2)
    ws2.append(["SO20261001001", "SKU-9001", 1, 299, "正常"])
    style_header(ws2, len(headers2))

    add_note_sheet(wb, "订单表模板 —— 导入到 orders + order_items 表", [
        "1. 第一个 sheet 是订单主表，第二个 sheet『订单明细』是每个订单买了什么。",
        "2. 订单号：唯一，重复导入按订单号覆盖更新。",
        "3. 订单状态：待付款 / 待发货 / 已发货 / 已签收 / 售后中 / 已退款 / 已取消。",
        "4. 时间格式：2026-10-01 10:00（签收时间尤其重要——AI 按它计算 7 天无理由/15 天质保等售后政策！）",
        "5. 快递公司/运单号/物流状态：有物流信息时 AI 才能回答查件问题。",
        "6. 明细里的 SKU编号 必须先存在于商品表中（先导入商品表，再导入订单表）。",
        "※ 数据行是示例，导入前请删除，或使用 --skip-sample 参数跳过。",
        "※ 保存后执行：python data/import_real_data.py 你的文件.xlsx --with-items",
    ])
    wb.save(OUT / "订单表模板.xlsx")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    make_products()
    make_orders()
    print(f"模板已生成:\n  {OUT / '商品表模板.xlsx'}\n  {OUT / '订单表模板.xlsx'}")


if __name__ == "__main__":
    main()
